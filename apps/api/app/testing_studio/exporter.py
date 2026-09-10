from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.testing_studio.constants import TestingType


TYPE_COLUMNS = {
    "FUNCTIONAL": ["Environment", "Browser / OS", "Step #", "Action", "Test Data", "Expected Result"],
    "API": ["Endpoint URL", "HTTP Method", "Headers JSON", "Request Body JSON", "Expected Status Code", "Expected Response JSON"],
    "DATABASE": ["Target Database", "Table Name", "Pre-requisite UI Action", "Verification SQL", "Expected Column", "Expected Value"],
    "AUTOMATION": ["Framework", "Script Identifier", "Git Repository", "Result Source", "Result Format", "Webhook Enabled"],
    "PERFORMANCE": ["Tool", "Virtual Users / Threads", "Ramp-up Seconds", "Duration Seconds", "Max Response Time ms", "Max Error Rate %", "Throughput"],
    "SECURITY": ["Vulnerability Category", "Target Vector", "Attack Payload", "Expected Defensive Behavior", "Severity", "CWE / CVE"],
    "ACCESSIBILITY": ["WCAG Clause", "Assistive Technology", "Alt Text", "Contrast Ratio", "Tab Order", "Keyboard Access", "Labels"],
}


def _string(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        import json
        return json.dumps(value, ensure_ascii=False, indent=2)
    return str(value)


class TestingStudioExporter:
    def generate(self, project, cases: list[dict], testing_type: str | None):
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        types = [testing_type] if testing_type else list(TYPE_COLUMNS)
        for type_value in types:
            if type_value not in TYPE_COLUMNS:
                continue
            ws = wb.create_sheet(title=type_value[:31])
            ws.append(["QABook", project.name, "Generated", __import__("datetime").datetime.now().strftime("%d-%b-%Y %H:%M")])
            ws.append([])
            headers = ["Test Case Code", "Requirement Code", "Scenario Code", "Module", "Priority", "Status", "Title", "Execution Method"] + TYPE_COLUMNS[type_value]
            ws.append(headers)
            for cell in ws[3]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for case in cases:
                profile = case.get("profile", {})
                attrs = profile.get("meta_attributes", {}) or {}
                if profile.get("testing_type") != type_value:
                    continue
                values = [case.get("test_case_code"), case.get("requirement_code", ""), case.get("scenario_code", ""), case.get("module"), case.get("priority"), case.get("status"), case.get("title"), profile.get("execution_method")]
                if type_value == "FUNCTIONAL":
                    steps = attrs.get("steps", []) or [{}]
                    for step in steps:
                        ws.append(values + [attrs.get("environment", ""), ", ".join(attrs.get("browser_os", [])), step.get("step_no", ""), step.get("action", ""), step.get("test_data", ""), step.get("expected_result", "")])
                elif type_value == "API":
                    values += [attrs.get("endpoint_url"), attrs.get("http_method"), _string(attrs.get("headers", {})), _string(attrs.get("request_body", {})), attrs.get("expected_status_code"), _string(attrs.get("expected_response", {}))]
                    ws.append(values)
                elif type_value == "DATABASE":
                    columns = attrs.get("expected_columns", []) or [{}]
                    for item in columns:
                        ws.append(values + [attrs.get("target_database"), attrs.get("table_name"), attrs.get("precondition_ui_action"), attrs.get("verification_sql"), item.get("column", ""), item.get("expected_value", "")])
                elif type_value == "AUTOMATION":
                    ws.append(values + [attrs.get("framework"), attrs.get("script_identifier"), attrs.get("git_repository"), attrs.get("result_source", "UPLOAD"), attrs.get("result_format", "JUnit XML"), attrs.get("webhook_enabled", False)])
                elif type_value == "PERFORMANCE":
                    sla = attrs.get("sla_benchmarks", {}) or {}
                    ws.append(values + [attrs.get("tool"), attrs.get("virtual_users"), attrs.get("ramp_up_seconds"), attrs.get("duration_seconds"), sla.get("max_response_time_ms"), sla.get("max_error_rate_percent"), sla.get("throughput")])
                elif type_value == "SECURITY":
                    ws.append(values + [attrs.get("vulnerability_category"), attrs.get("target_vector"), attrs.get("attack_payload"), attrs.get("expected_defensive_behavior"), attrs.get("severity", ""), attrs.get("cwe_cve", "")])
                elif type_value == "ACCESSIBILITY":
                    flags = attrs.get("audit_flags", {}) or {}
                    ws.append(values + [attrs.get("wcag_clause"), attrs.get("assistive_technology"), flags.get("alt_text", False), flags.get("contrast_ratio", False), flags.get("tab_order", False), flags.get("keyboard_access", False), flags.get("labels", False)])
            for column in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(column)].width = min(max(14, max(len(str(ws.cell(row=r, column=column).value or "")) for r in range(1, min(ws.max_row, 30) + 1)) + 2), 45)
            ws.freeze_panes = "A4"
            ws.auto_filter.ref = f"A3:{get_column_letter(ws.max_column)}{ws.max_row}"
            for row in ws.iter_rows(min_row=4):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
