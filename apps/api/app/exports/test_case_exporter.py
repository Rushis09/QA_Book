from datetime import datetime
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from app.exports.constants import ExportConstants


TYPE_COLUMNS = {
    "FUNCTIONAL": [
        "Environment",
        "Browser / OS",
        "Step #",
        "Action",
        "Step Test Data",
        "Step Expected Result",
    ],
    "API": [
        "Authentication",
        "Endpoint URL",
        "HTTP Method",
        "Headers JSON",
        "Path Parameters JSON",
        "Request Body JSON",
        "Expected Status Code",
        "Expected Response JSON",
    ],
    "DATABASE": [
        "Target Database",
        "Schema Name",
        "Table Name",
        "Pre-requisite UI Action",
        "Verification SQL",
        "Expected Column",
        "Expected Value",
    ],
    "AUTOMATION": [
        "Framework",
        "Script Identifier",
        "Git Repository",
        "Result Source",
        "Result Format",
        "Webhook Enabled",
    ],
    "PERFORMANCE": [
        "Performance Model",
        "Tool",
        "Virtual Users / Threads",
        "Ramp-up Seconds",
        "Arrival Rate",
        "Duration Seconds",
        "Target Endpoint / Operation",
        "Workload Description",
        "Max Response Time ms",
        "Max Error Rate %",
        "Throughput",
    ],
    "SECURITY": [
        "Vulnerability Category",
        "Target Vector",
        "Attack Payload",
        "Expected Defensive Behavior",
        "Severity",
        "CWE",
        "CVE",
    ],
    "ACCESSIBILITY": [
        "WCAG Clause",
        "Assistive Technology",
        "Page / Component",
        "Expected Accessible Behavior",
        "Alt Text",
        "Contrast Ratio",
        "Tab Order",
        "Keyboard Access",
        "Labels",
    ],
}


def _string(value):
    """Convert values to Excel-friendly strings."""

    if value is None:
        return ""

    if isinstance(value, (dict, list)):
        return json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
        )

    return str(value)


def _profile_values(test_case):
    """
    Return normalized testing profile information.

    Legacy Test Cases without a profile are treated as
    Functional + Manual.
    """

    profile = test_case.get("profile") or {}

    testing_type = profile.get(
        "testing_type",
        "FUNCTIONAL",
    )

    execution_method = profile.get(
        "execution_method",
        "MANUAL",
    )

    meta_attributes = profile.get(
        "meta_attributes",
        {},
    ) or {}

    return (
        testing_type,
        execution_method,
        meta_attributes,
    )


def _get_first(attrs, *keys, default=""):
    """Return the first available value from multiple possible keys."""

    for key in keys:
        value = attrs.get(key)

        if value is not None and value != "":
            return value

    return default


class TestCaseExporter:
    """
    Canonical Excel exporter for Test Cases and Testing Studio.

    The same exporter is used by:
    - Normal Test Case Export
    - Testing Studio Export
    """

    def generate(
        self,
        metadata: dict,
        test_cases: list[dict],
    ) -> BytesIO:
        """Generate the complete type-aware Test Case workbook."""

        workbook = Workbook()

        default_sheet = workbook.active
        workbook.remove(default_sheet)

        for testing_type, type_columns in TYPE_COLUMNS.items():

            worksheet = workbook.create_sheet(
                title=testing_type[:31],
            )

            common_columns = [
                "Test Case Code",
                "Requirement Code",
                "Scenario Code",
                "Module",
                "Testing Type",
                "Execution Method",
                "Priority",
                "Status",
                "Title",
                "Description",
                "Automation Eligibility",
                "Automation Status",
                "Preconditions",
                "Test Data",
                "Expected Result",
            ]

            headers = common_columns + type_columns

            header_row = 11

            self._add_document_header(
                worksheet=worksheet,
                metadata=metadata,
                total_columns=len(headers),
            )

            self._add_table_header(
                worksheet=worksheet,
                headers=headers,
                header_row=header_row,
            )

            for test_case in test_cases:

                (
                    case_type,
                    execution_method,
                    attrs,
                ) = _profile_values(test_case)

                if case_type != testing_type:
                    continue

                common_values = [
                    test_case.get(
                        "test_case_code",
                        "",
                    ),
                    test_case.get(
                        "requirement_code",
                        "",
                    ),
                    test_case.get(
                        "scenario_code",
                        "",
                    ),
                    test_case.get(
                        "module",
                        "",
                    ),
                    case_type,
                    execution_method,
                    test_case.get(
                        "priority",
                        "",
                    ),
                    test_case.get(
                        "status",
                        "",
                    ),
                    test_case.get(
                        "title",
                        "",
                    ),
                    test_case.get(
                        "description",
                        "",
                    ),
                    test_case.get(
                        "automation_eligibility",
                        "",
                    ),
                    test_case.get(
                        "automation_status",
                        "",
                    ),
                    test_case.get(
                        "preconditions",
                        "",
                    ),
                    test_case.get(
                        "test_data",
                        "",
                    ),
                    test_case.get(
                        "expected_result",
                        "",
                    ),
                ]

                if testing_type == "FUNCTIONAL":

                    self._write_functional(
                        worksheet=worksheet,
                        common_values=common_values,
                        test_case=test_case,
                        attrs=attrs,
                    )

                elif testing_type == "API":

                    worksheet.append(
                        common_values
                        + [
                            _get_first(
                                attrs,
                                "authentication",
                                "auth",
                            ),
                            _get_first(
                                attrs,
                                "endpoint_url",
                            ),
                            _get_first(
                                attrs,
                                "http_method",
                            ),
                            _string(
                                attrs.get(
                                    "headers",
                                    {},
                                )
                            ),
                            _string(
                                _get_first(
                                    attrs,
                                    "path_parameters",
                                    "path_params",
                                    default={},
                                )
                            ),
                            _string(
                                attrs.get(
                                    "request_body",
                                    {},
                                )
                            ),
                            _get_first(
                                attrs,
                                "expected_status_code",
                            ),
                            _string(
                                attrs.get(
                                    "expected_response",
                                    {},
                                )
                            ),
                        ]
                    )

                elif testing_type == "DATABASE":

                    self._write_database(
                        worksheet=worksheet,
                        common_values=common_values,
                        attrs=attrs,
                    )

                elif testing_type == "AUTOMATION":

                    worksheet.append(
                        common_values
                        + [
                            attrs.get(
                                "framework",
                                "",
                            ),
                            attrs.get(
                                "script_identifier",
                                "",
                            ),
                            attrs.get(
                                "git_repository",
                                "",
                            ),
                            attrs.get(
                                "result_source",
                                "UPLOAD",
                            ),
                            attrs.get(
                                "result_format",
                                "JUnit XML",
                            ),
                            attrs.get(
                                "webhook_enabled",
                                False,
                            ),
                        ]
                    )

                elif testing_type == "PERFORMANCE":

                    self._write_performance(
                        worksheet=worksheet,
                        common_values=common_values,
                        attrs=attrs,
                    )

                elif testing_type == "SECURITY":

                    self._write_security(
                        worksheet=worksheet,
                        common_values=common_values,
                        attrs=attrs,
                    )

                elif testing_type == "ACCESSIBILITY":

                    self._write_accessibility(
                        worksheet=worksheet,
                        common_values=common_values,
                        attrs=attrs,
                    )

            self._format_worksheet(
                worksheet=worksheet,
                header_row=header_row,
            )

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return output

    @staticmethod
    def _add_document_header(
        worksheet,
        metadata,
        total_columns,
    ):
        """Add the standard QABook document header."""

        last_column = get_column_letter(
            total_columns,
        )

        # Row 1 - QABook
        worksheet.merge_cells(
            f"A1:{last_column}1"
        )

        title_cell = worksheet["A1"]

        title_cell.value = "QABook"

        title_cell.font = Font(
            bold=True,
            size=16,
        )

        title_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[1].height = 24

        # Row 2 - Document title
        worksheet.merge_cells(
            f"A2:{last_column}2"
        )

        document_cell = worksheet["A2"]

        document_cell.value = (
            ExportConstants.TEST_CASES_DOCUMENT_TITLE
        )

        document_cell.font = Font(
            bold=True,
            size=11,
        )

        document_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        worksheet.row_dimensions[2].height = 20

        # Row 4 - Project Information
        worksheet.merge_cells(
            f"A4:{last_column}4"
        )

        project_information_cell = worksheet["A4"]

        project_information_cell.value = (
            "Project Information"
        )

        project_information_cell.font = Font(
            bold=True,
        )

        project_information_cell.fill = PatternFill(
            "solid",
            fgColor="E7E6E6",
        )

        project_information_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        worksheet.row_dimensions[4].height = 20

        project_name = metadata.get(
            ExportConstants.PROJECT_NAME_LABEL,
            metadata.get(
                "project_name",
                "",
            ),
        )

        project_code = metadata.get(
            ExportConstants.PROJECT_CODE_LABEL,
            metadata.get(
                "project_code",
                "",
            ),
        )

        generated_by = metadata.get(
            ExportConstants.GENERATED_BY_LABEL,
            "QABook",
        )

        generated_date = metadata.get(
            ExportConstants.GENERATED_DATE_LABEL,
            datetime.now().strftime(
                "%d-%b-%Y %H:%M"
            ),
        )

        version = metadata.get(
            ExportConstants.VERSION_LABEL,
            ExportConstants.QDS_VERSION,
        )

        information = [
            (
                5,
                "Project Name",
                project_name,
            ),
            (
                6,
                "Project Code",
                project_code,
            ),
            (
                7,
                "Generated By",
                generated_by,
            ),
            (
                8,
                "Generated Date",
                generated_date,
            ),
            (
                9,
                "Version",
                version,
            ),
        ]

        thin_side = Side(
            style="thin",
            color="D9D9D9",
        )

        information_border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for row, label, value in information:

            label_cell = worksheet.cell(
                row=row,
                column=1,
            )

            label_cell.value = label

            label_cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            label_cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

            label_cell.border = information_border

            label_cell.alignment = Alignment(
                vertical="center",
            )

            value_cell = worksheet.cell(
                row=row,
                column=2,
            )

            value_cell.value = value

            value_cell.border = information_border

            value_cell.alignment = Alignment(
                vertical="center",
                wrap_text=True,
            )

    @staticmethod
    def _add_table_header(
        worksheet,
        headers,
        header_row,
    ):
        """Create and style the main Test Case table header."""

        thin_side = Side(
            style="thin",
            color="D9D9D9",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for column, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=header_row,
                column=column,
            )

            cell.value = header

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        worksheet.row_dimensions[
            header_row
        ].height = 32

    @staticmethod
    def _write_functional(
        worksheet,
        common_values,
        test_case,
        attrs,
    ):
        """Write Functional Test Case rows."""

        steps = attrs.get(
            "steps",
            [],
        ) or []

        if not steps:
            steps = _legacy_steps(
                test_case,
            )

        # A Functional Test Case without steps still
        # gets one row so its common fields are exported.
        if not steps:
            steps = [{}]

        legacy_test_data = (
            attrs.get(
                "legacy_test_data",
            )
            or test_case.get(
                "test_data",
                "",
            )
        )

        legacy_expected_result = (
            attrs.get(
                "legacy_expected_result",
            )
            or test_case.get(
                "expected_result",
                "",
            )
        )

        browser_os = attrs.get(
            "browser_os",
            [],
        )

        if isinstance(browser_os, list):
            browser_os = ", ".join(
                str(item)
                for item in browser_os
            )

        for step in steps:

            step = step or {}

            step_test_data = (
                step.get(
                    "test_data",
                )
                or legacy_test_data
            )

            step_expected_result = (
                step.get(
                    "expected_result",
                )
                or legacy_expected_result
            )

            worksheet.append(
                common_values
                + [
                    attrs.get(
                        "environment",
                        "",
                    ),
                    browser_os,
                    step.get(
                        "step_no",
                        "",
                    ),
                    step.get(
                        "action",
                        "",
                    ),
                    step_test_data,
                    step_expected_result,
                ]
            )

    @staticmethod
    def _write_database(
        worksheet,
        common_values,
        attrs,
    ):
        """Write Database Test Case rows."""

        columns = (
            attrs.get(
                "expected_columns",
            )
            or []
        )

        if not columns:
            columns = [{}]

        for item in columns:

            item = item or {}

            worksheet.append(
                common_values
                + [
                    attrs.get(
                        "target_database",
                        "",
                    ),
                    attrs.get(
                        "schema_name",
                        "",
                    ),
                    attrs.get(
                        "table_name",
                        "",
                    ),
                    attrs.get(
                        "precondition_ui_action",
                        "",
                    ),
                    attrs.get(
                        "verification_sql",
                        "",
                    ),
                    item.get(
                        "column",
                        "",
                    ),
                    item.get(
                        "expected_value",
                        "",
                    ),
                ]
            )

    @staticmethod
    def _write_performance(
        worksheet,
        common_values,
        attrs,
    ):
        """Write Performance Test Case row."""

        sla = (
            attrs.get(
                "sla_benchmarks",
                {},
            )
            or {}
        )

        worksheet.append(
            common_values
            + [
                _get_first(
                    attrs,
                    "performance_model",
                    "model",
                ),
                attrs.get(
                    "tool",
                    "",
                ),
                attrs.get(
                    "virtual_users",
                    "",
                ),
                attrs.get(
                    "ramp_up_seconds",
                    "",
                ),
                attrs.get(
                    "arrival_rate",
                    "",
                ),
                attrs.get(
                    "duration_seconds",
                    "",
                ),
                _get_first(
                    attrs,
                    "target_endpoint_or_operation",
                    "target_endpoint",
                    "target_operation",
                ),
                attrs.get(
                    "workload_description",
                    "",
                ),
                sla.get(
                    "max_response_time_ms",
                    "",
                ),
                sla.get(
                    "max_error_rate_percent",
                    "",
                ),
                sla.get(
                    "throughput",
                    "",
                ),
            ]
        )

    @staticmethod
    def _write_security(
        worksheet,
        common_values,
        attrs,
    ):
        """Write Security Test Case row."""

        combined_cwe_cve = attrs.get(
            "cwe_cve",
            "",
        )

        cwe = attrs.get(
            "cwe",
            "",
        )

        cve = attrs.get(
            "cve",
            "",
        )

        if not cwe and not cve and combined_cwe_cve:
            cwe = combined_cwe_cve

        worksheet.append(
            common_values
            + [
                attrs.get(
                    "vulnerability_category",
                    "",
                ),
                attrs.get(
                    "target_vector",
                    "",
                ),
                attrs.get(
                    "attack_payload",
                    "",
                ),
                attrs.get(
                    "expected_defensive_behavior",
                    "",
                ),
                attrs.get(
                    "severity",
                    "",
                ),
                cwe,
                cve,
            ]
        )

    @staticmethod
    def _write_accessibility(
        worksheet,
        common_values,
        attrs,
    ):
        """Write Accessibility Test Case row."""

        flags = (
            attrs.get(
                "audit_flags",
                {},
            )
            or {}
        )

        worksheet.append(
            common_values
            + [
                attrs.get(
                    "wcag_clause",
                    "",
                ),
                attrs.get(
                    "assistive_technology",
                    "",
                ),
                _get_first(
                    attrs,
                    "page_component",
                    "page_or_component",
                ),
                _get_first(
                    attrs,
                    "expected_accessible_behavior",
                    "expected_behavior",
                ),
                flags.get(
                    "alt_text",
                    False,
                ),
                flags.get(
                    "contrast_ratio",
                    False,
                ),
                flags.get(
                    "tab_order",
                    False,
                ),
                flags.get(
                    "keyboard_access",
                    False,
                ),
                flags.get(
                    "labels",
                    False,
                ),
            ]
        )

    @staticmethod
    def _format_worksheet(
        worksheet,
        header_row,
    ):
        """Apply common worksheet formatting."""

        # Freeze below the table header.
        worksheet.freeze_panes = (
            f"A{header_row + 1}"
        )

        # Filter starts from the actual table header.
        worksheet.auto_filter.ref = (
            f"A{header_row}:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )

        # Body formatting.
        for row in worksheet.iter_rows(
            min_row=header_row + 1,
        ):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        # Automatic column widths.
        for column in range(
            1,
            worksheet.max_column + 1,
        ):
            max_length = max(
                len(
                    str(
                        worksheet.cell(
                            row=row,
                            column=column,
                        ).value
                        or ""
                    )
                )
                for row in range(
                    1,
                    min(
                        worksheet.max_row,
                        30,
                    )
                    + 1,
                )
            )

            worksheet.column_dimensions[
                get_column_letter(column)
            ].width = min(
                max(
                    14,
                    max_length + 2,
                ),
                45,
            )

        # Body borders.
        thin_side = Side(
            style="thin",
            color="D9D9D9",
        )

        border = Border(
            left=thin_side,
            right=thin_side,
            top=thin_side,
            bottom=thin_side,
        )

        for row in worksheet.iter_rows(
            min_row=header_row + 1,
        ):
            for cell in row:
                cell.border = border


def _legacy_steps(test_case):
    """Convert legacy multiline steps into structured steps."""

    raw_steps = test_case.get("steps") or ""

    if not raw_steps:
        return []

    lines = [
        line.strip()
        for line in str(raw_steps).splitlines()
        if line.strip()
    ]

    return [
        {
            "step_no": index,
            "action": line,
            "test_data": "",
            "expected_result": "",
        }
        for index, line in enumerate(
            lines,
            start=1,
        )
    ]